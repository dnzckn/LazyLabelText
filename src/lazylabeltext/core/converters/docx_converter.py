"""DOCX document converter using python-docx."""

from __future__ import annotations

from pathlib import Path

from lazylabeltext.core.converters import _register
from lazylabeltext.core.exceptions import DocumentParseError
from lazylabeltext.core.models import ConvertedDocument, Heading, SectionBoundary


class DOCXConverter:
    """Converts DOCX files to internal representation."""

    def convert(self, file_path: str) -> ConvertedDocument:
        path = Path(file_path)
        try:
            from docx import Document
        except ImportError as e:
            raise DocumentParseError(file_path, "python-docx not installed") from e

        try:
            doc = Document(str(path))
        except Exception as e:
            raise DocumentParseError(file_path, str(e)) from e

        text_parts: list[str] = []
        headings: list[Heading] = []
        char_offset = 0

        for para in doc.paragraphs:
            para_text = para.text
            if not para_text.strip():
                text_parts.append("")
                char_offset += 1  # newline
                continue

            style_name = para.style.name if para.style else ""
            heading_level = self._parse_heading_level(style_name)

            if heading_level is not None:
                headings.append(
                    Heading(
                        level=heading_level,
                        text=para_text.strip(),
                        char_start=char_offset,
                        char_end=char_offset + len(para_text),
                    )
                )

            text_parts.append(para_text)
            char_offset += len(para_text) + 1  # +1 for newline

        full_text = "\n".join(text_parts)

        # Build sections from headings
        sections = self._build_sections(headings, len(full_text))

        # Extract metadata
        metadata: dict = {"source": str(path)}
        core = doc.core_properties
        if core.author:
            metadata["author"] = core.author
        if core.title:
            metadata["title"] = core.title
        if core.created:
            metadata["created"] = str(core.created)

        return ConvertedDocument(
            filename=path.name,
            format="docx",
            full_text=full_text,
            headings=headings,
            pages=[],
            sections=sections,
            metadata=metadata,
            status="parsed",
        )

    def _parse_heading_level(self, style_name: str) -> int | None:
        """Extract heading level from DOCX style name."""
        if not style_name:
            return None
        style_lower = style_name.lower()
        if style_lower.startswith("heading"):
            parts = style_lower.replace("heading", "").strip()
            try:
                return int(parts)
            except ValueError:
                return None
        return None

    def _build_sections(
        self, headings: list[Heading], text_length: int
    ) -> list[SectionBoundary]:
        if not headings:
            return [
                SectionBoundary(section_path=[], char_start=0, char_end=text_length)
            ]

        sections = []
        path_stack: list[str] = []

        for i, heading in enumerate(headings):
            while len(path_stack) >= heading.level:
                path_stack.pop()
            path_stack.append(heading.text)

            end = headings[i + 1].char_start if i + 1 < len(headings) else text_length
            sections.append(
                SectionBoundary(
                    section_path=list(path_stack),
                    char_start=heading.char_start,
                    char_end=end,
                )
            )
        return sections

    def supported_extensions(self) -> set[str]:
        return {".docx"}


_register({".docx"}, DOCXConverter())
